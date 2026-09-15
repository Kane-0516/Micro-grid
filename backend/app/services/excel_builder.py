"""Build the solution Excel workbook."""

from __future__ import annotations

import datetime
import io
import math
from dataclasses import dataclass
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

NAVY = "003C71"
SUBHDR = "1F4E79"
LIGHT_BLUE = "DEEAF1"
LIGHT_GREEN = "E2EFDA"
LIGHT_GRAY = "F2F2F2"
WHITE = "FFFFFF"
ACCENT = "FFC000"
HORIZONTAL_CENTER = "center"
VERTICAL_CENTER = "center"
DIESEL_GENERATOR_LABEL = "Diesel Generator"
ECONOMIC_HIGHLIGHTS = {"Payback Period", "20-year Net Revenue"}
COMPARISON_HEADERS = (
    "Year",
    "MG Annual",
    "Diesel Annual",
    "MG Cumulative",
    "Diesel Cumulative",
    "MG LCOE",
    "Diesel LCOE",
    "Annual Revenue",
    "Cumulative Revenue",
)


@dataclass(frozen=True)
class ExcelBuildContext:
    """Bundled solution data used to fill in the Excel report."""

    sc: dict[str, Any]
    cx: dict[str, Any]
    sim: dict[str, Any]
    sm: dict[str, Any]
    ct: list[Any]
    con: Any
    now: str
    pv_sets: int
    pv_pps: int
    pv_total: int
    pv_kw: float
    pv_model: str
    pv_watts: int
    bat_count: int
    bat_kwh_total: float
    bat_kwh_each: float
    bat_model: str
    gen_kw: float
    gen_model: str
    volt_level: str
    ems_mode: str
    ann_load: float
    area_m2: float
    num_inv: int
    layout_note: str
    measured_site_area: str
    usable_site_area: str
    layout_max_sets: Any
    project_type: str
    pv_mod_cost: float
    mount_cost: float
    bess_cost: float
    gen_cost: float
    trans_cost: float
    inst_cost: float
    acc_cost: float
    other_cost: float
    subtotal: float
    profit_pct: float
    profit_amt: float
    selling: float


def _parse_excel_context(req) -> ExcelBuildContext:
    sc = req.systemConfig or {}
    cx = req.capex or {}
    sim = req.simulation or {}
    sm = req.summary or {}
    ct = req.comparisonTable or []

    pv_sets = int(sc.get("bracketSets", 0))
    pv_pps = int(sc.get("panelsPerSet", 32))
    pv_total = pv_sets * pv_pps
    pv_kw = float(sc.get("pvCapacityKw", 0))
    pv_model = str(sc.get("panelModel", "-"))
    pv_watts = int(sc.get("panelWatts", 655))
    bat_count = int(sc.get("batteryPackCount", 0))
    bat_kwh_total = float(sc.get("batteryCapacityKwh", 0))
    bat_kwh_each = round(bat_kwh_total / bat_count, 1) if bat_count else 16
    bat_model = str(sc.get("batteryModel", "LFP-16kWh"))
    gen_kw = float(sc.get("dieselCapacityKw", 0))
    gen_model = str(sc.get("dieselModel", "-"))

    return ExcelBuildContext(
        sc=sc,
        cx=cx,
        sim=sim,
        sm=sm,
        ct=ct,
        con=req.contact,
        now=datetime.datetime.now().strftime("%Y-%m-%d"),
        pv_sets=pv_sets,
        pv_pps=pv_pps,
        pv_total=pv_total,
        pv_kw=pv_kw,
        pv_model=pv_model,
        pv_watts=pv_watts,
        bat_count=bat_count,
        bat_kwh_total=bat_kwh_total,
        bat_kwh_each=bat_kwh_each,
        bat_model=bat_model,
        gen_kw=gen_kw,
        gen_model=gen_model,
        volt_level=str(sc.get("voltageLevel", "48V")),
        ems_mode=str(sc.get("emsMode", "edge")),
        ann_load=float(sc.get("annualLoadKwh", 0)),
        area_m2=float(sc.get("occupiedAreaM2", 0)),
        num_inv=max(1, math.ceil(bat_count / 6)),
        layout_note=str(sm.get("siteLayoutNote", "") or "").strip(),
        measured_site_area=str(
            sm.get("measuredSiteAreaDisplay", "") or ""
        ).strip(),
        usable_site_area=str(sm.get("usableSiteAreaDisplay", "") or "").strip(),
        layout_max_sets=sm.get("layoutMaxBracketSets"),
        project_type=_project_type(pv_kw, gen_kw),
        pv_mod_cost=float(cx.get("pvModuleCost", 0)),
        mount_cost=float(cx.get("pvMountingCost", 0)),
        bess_cost=float(cx.get("energyStorageCost", 0)),
        gen_cost=float(cx.get("dieselGeneratorCost", 0)),
        trans_cost=float(cx.get("intlTransportCost", 0)),
        inst_cost=float(cx.get("installationCost", 0)),
        acc_cost=float(cx.get("accessoryCost", 0)),
        other_cost=float(cx.get("otherInitialCost", 0)),
        subtotal=float(cx.get("equipmentSubtotal", 0)),
        profit_pct=float(cx.get("profitMargin", 0)),
        profit_amt=float(cx.get("profitAmount", 0)),
        selling=float(cx.get("sellingPrice", 0)),
    )


def _build_bom_rows(ctx: ExcelBuildContext) -> list[tuple]:
    pv_dimensions = "2384x1303x33 mm" if ctx.pv_watts >= 655 else "-"
    bom_rows = [
        (
            1,
            "",
            "Hybrid Inverter",
            "Hybrid Inverter",
            "18K-2P-LV",
            "",
            ctx.num_inv,
            "EA",
            f"Input: PV + Battery ({ctx.volt_level}), Output: AC 230V/50Hz; "
            "supports off-grid & grid-tied modes",
        ),
        (
            2,
            "",
            "Battery Pack",
            "Battery Pack",
            ctx.bat_model,
            "",
            ctx.bat_count,
            "EA",
            f"LFP, {ctx.bat_kwh_each} kWh/pack, "
            f"total {ctx.bat_kwh_total:.0f} kWh; "
            "cycle life >= 4,000 cycles; BMS included",
        ),
        (
            3,
            "",
            "Industrial PC / EMS",
            "Industrial PC / EMS",
            "ECU-1170",
            "",
            1,
            "EA",
            f"VoltageEnergy EMS; control mode: {ctx.ems_mode}; "
            "real-time monitoring and remote O&M",
        ),
        (
            4,
            "ES1-S002-1",
            "Integrated Skid",
            "Integrated Skid",
            "",
            "",
            ctx.pv_sets,
            "SET",
            "Main material: Q355B hot-dip galvanized; integrated pallet "
            "spliced weldment",
        ),
        (
            5,
            "ES1-S001-1",
            "Foldable Mounting System",
            "Foldable Mounting System",
            "",
            "",
            ctx.pv_sets,
            "SET",
            f"A set consisting of {ctx.pv_pps} rows of PV modules; "
            "steel-aluminum structure; foldable design",
        ),
        (
            6,
            "",
            "BOS Box",
            "BOS Box",
            "",
            "",
            ctx.pv_sets,
            "EA",
            "Balance of System box; includes DC/AC breakers, combiner, surge "
            "protection, metering",
        ),
        (
            7,
            "ES1-E004-1",
            "PV Panel Module",
            "PV Panel Module",
            ctx.pv_model,
            pv_dimensions,
            ctx.pv_total,
            "EA",
            f"{ctx.pv_model}; {ctx.pv_watts}Wp; "
            f"total {ctx.pv_kw:.1f} kW; "
            f"footprint about {ctx.area_m2:.0f} m2",
        ),
    ]
    _append_diesel_bom_row(bom_rows, ctx.sim, ctx.gen_kw, ctx.gen_model)
    return bom_rows


def _cfont(
    size: int = 11,
    bold: bool = False,
    color: str = "000000",
    italic: bool = False,
) -> Font:
    return Font(
        name="Calibri", size=size, bold=bold, color=color, italic=italic
    )


def _sfill(hex_color: str) -> PatternFill:
    return PatternFill("solid", fgColor=hex_color)


def _tborder(style: str = "thin") -> Border:
    side = Side(style=style, color="8EA9C1")
    return Border(left=side, right=side, top=side, bottom=side)


def _cell_set(cell, value, font=None, fill_=None, align=None, border=None):
    cell.value = value
    if font:
        cell.font = font
    if fill_:
        cell.fill = fill_
    if align:
        cell.alignment = align
    if border:
        cell.border = border


def _set_col_w(ws, widths):
    for index, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(index)].width = width


def _fmt_money(value: float) -> str:
    return f"$ {value:,.0f}"


def _fmt_optional_money(value: float) -> str:
    return _fmt_money(value) if value else "-"


def _sheet_title(ws, cell_ref: str, text: str):
    ws.merge_cells(f"{cell_ref}:B1")
    ws.row_dimensions[1].height = 26
    _cell_set(
        ws[cell_ref],
        text,
        font=_cfont(13, bold=True, color="FFFFFF"),
        fill_=_sfill(SUBHDR),
        align=Alignment(horizontal="center", vertical="center"),
    )


def _project_type(pv_kw: float, gen_kw: float) -> str:
    if pv_kw <= 0:
        return "Off-grid Power System"
    return (
        "PV + BESS + Diesel Microgrid" if gen_kw > 0 else "PV + BESS Microgrid"
    )


def _layout_note_lines(
    layout_note: str,
    measured_site_area: str,
    usable_site_area: str,
    layout_max_sets,
) -> list[str]:
    lines = [layout_note]
    optional_lines = (
        (measured_site_area, f"Measured site area: {measured_site_area}"),
        (usable_site_area, f"Area basis used for layout: {usable_site_area}"),
        (
            layout_max_sets not in (None, ""),
            f"Layout-based max bracket sets: {layout_max_sets}",
        ),
    )
    lines.extend(text for value, text in optional_lines if value)
    return lines


def _style_table_row(
    ws,
    row_index: int,
    values,
    *,
    fill_color: str,
    height: int,
    font_factory,
    alignment_factory,
) -> None:
    for col_index, value in enumerate(values, 1):
        cell = ws.cell(row_index, col_index, value)
        cell.font = font_factory(col_index)
        cell.fill = _sfill(fill_color)
        cell.border = _tborder()
        cell.alignment = alignment_factory(col_index)
    ws.row_dimensions[row_index].height = height


def _alternating_color(index: int, highlight: bool = False) -> str:
    if highlight:
        return LIGHT_GREEN
    return LIGHT_GRAY if index % 2 == 0 else WHITE


def _write_economic_rows(ws, start_row: int, rows) -> int:
    current_row = start_row
    for label, value in rows:
        highlight = label in ECONOMIC_HIGHLIGHTS
        _style_table_row(
            ws,
            current_row,
            (label, value),
            fill_color=_alternating_color(current_row, highlight),
            height=18,
            font_factory=lambda _column, bold=highlight: _cfont(11, bold=bold),
            alignment_factory=lambda column: Alignment(
                horizontal="right" if column == 2 else "left",
                vertical=VERTICAL_CENTER,
            ),
        )
        current_row += 1
    return current_row


def _comparison_row_values(row) -> list:
    return [
        row.get("year"),
        f"${row.get('mgAnnualCost', 0):,.0f}",
        f"${row.get('dieselAnnualCost', 0):,.0f}",
        f"${row.get('mgCumulative', 0):,.0f}",
        f"${row.get('dieselCumulative', 0):,.0f}",
        f"${row.get('mgLcoe', 0):.4f}",
        f"${row.get('dieselLcoe', 0):.4f}",
        f"${row.get('annualRevenue', 0):,.0f}",
        f"${row.get('cumulativeRevenue', 0):,.0f}",
    ]


def _write_comparison_table(ws, start_row: int, rows, breakeven_year) -> int:
    if not rows:
        return start_row
    current_row = start_row
    ws.merge_cells(f"A{current_row}:I{current_row}")
    ws.row_dimensions[current_row].height = 24
    _cell_set(
        ws[f"A{current_row}"],
        "Annual Cost Comparison (MG vs Diesel-only)",
        font=_cfont(12, bold=True, color=WHITE),
        fill_=_sfill(SUBHDR),
        align=Alignment(horizontal=HORIZONTAL_CENTER, vertical=VERTICAL_CENTER),
    )
    current_row += 1
    _style_table_row(
        ws,
        current_row,
        COMPARISON_HEADERS,
        fill_color="2E74B5",
        height=34,
        font_factory=lambda _column: _cfont(10, bold=True, color=WHITE),
        alignment_factory=lambda _column: Alignment(
            horizontal=HORIZONTAL_CENTER,
            vertical=VERTICAL_CENTER,
            wrap_text=True,
        ),
    )
    current_row += 1
    for index, row in enumerate(rows):
        is_breakeven = (
            breakeven_year is not None and row.get("year") == breakeven_year
        )
        _style_table_row(
            ws,
            current_row,
            _comparison_row_values(row),
            fill_color=_alternating_color(index, is_breakeven),
            height=16,
            font_factory=lambda _column, bold=is_breakeven: _cfont(
                10, bold=bold
            ),
            alignment_factory=lambda _column: Alignment(
                horizontal=HORIZONTAL_CENTER, vertical=VERTICAL_CENTER
            ),
        )
        current_row += 1
    return current_row


def _write_bom_rows(ws, rows) -> None:
    for offset, row_data in enumerate(rows, start=4):
        _style_table_row(
            ws,
            offset,
            row_data,
            fill_color=_alternating_color(offset),
            height=42,
            font_factory=lambda column: _cfont(11, bold=(column in (3, 4))),
            alignment_factory=lambda column: Alignment(
                horizontal=HORIZONTAL_CENTER
                if column in (1, 6, 7, 8)
                else "left",
                vertical=VERTICAL_CENTER,
                wrap_text=True,
            ),
        )


def _append_diesel_bom_row(
    rows, simulation, gen_kw: float, gen_model: str
) -> None:
    if gen_kw <= 0:
        return
    diesel_hours = float(simulation.get("mgDieselHours", 0))
    diesel_only_hours = max(float(simulation.get("dieselRunHoursA", 8760)), 1)
    reduction = (1 - diesel_hours / diesel_only_hours) * 100
    model = (
        gen_model
        if gen_model and gen_model != "-"
        else f"{gen_kw:.0f}kW genset"
    )
    rows.append(
        (
            8,
            "ES1-E005-1",
            DIESEL_GENERATOR_LABEL,
            DIESEL_GENERATOR_LABEL,
            model,
            "",
            1,
            "EA",
            f"{gen_kw:.0f} kW; backup power for cloudy days; "
            f"annual run hours reduced about {reduction:.0f}% "
            "vs diesel-only",
        )
    )


def _write_capex_rows(ws, start_row: int, rows) -> int:
    current_row = start_row
    for label, amount in rows:
        _style_table_row(
            ws,
            current_row,
            (label, _fmt_optional_money(amount)),
            fill_color=_alternating_color(current_row),
            height=18,
            font_factory=lambda column: _cfont(11, bold=(column == 1)),
            alignment_factory=lambda column: Alignment(
                horizontal="right" if column == 2 else "left",
                vertical=VERTICAL_CENTER,
            ),
        )
        current_row += 1
    return current_row


def _write_pricing_rows(ws, start_row: int, rows) -> int:
    current_row = start_row
    for label, value, fill_color in rows:
        quote_row = "Quote" in label
        _style_table_row(
            ws,
            current_row,
            (label, value),
            fill_color=fill_color,
            height=20,
            font_factory=lambda _column, quote=quote_row: _cfont(
                11, bold=True, color=NAVY if quote else "000000"
            ),
            alignment_factory=lambda column: Alignment(
                horizontal="right" if column == 2 else "left",
                vertical=VERTICAL_CENTER,
            ),
        )
        current_row += 1
    return current_row


def _write_site_layout_note(
    ws,
    current_row: int,
    layout_note: str,
    measured_site_area: str,
    usable_site_area: str,
    layout_max_sets,
) -> int:
    if not layout_note:
        return current_row
    ws.merge_cells(f"A{current_row}:B{current_row}")
    ws.row_dimensions[current_row].height = 48
    layout_lines = _layout_note_lines(
        layout_note, measured_site_area, usable_site_area, layout_max_sets
    )
    _cell_set(
        ws[f"A{current_row}"],
        "Site Layout Note\n" + "\n".join(layout_lines),
        font=_cfont(10, color="404040"),
        fill_=_sfill(LIGHT_BLUE),
        align=Alignment(
            horizontal="left", vertical=VERTICAL_CENTER, wrap_text=True
        ),
        border=_tborder(),
    )
    return current_row + 2


def _add_contact_sheet(workbook, contact, report_date: str) -> None:
    ws = workbook.create_sheet("Contact")
    _set_col_w(ws, [28, 40])
    ws.merge_cells("A1:B1")
    ws.row_dimensions[1].height = 26
    _cell_set(
        ws["A1"],
        "Customer Contact Information",
        font=_cfont(13, bold=True, color=WHITE),
        fill_=_sfill(SUBHDR),
        align=Alignment(horizontal=HORIZONTAL_CENTER, vertical=VERTICAL_CENTER),
    )
    rows = [
        ("Full Name", f"{contact.firstName} {contact.lastName}".strip()),
        ("Company", contact.company),
        ("Email", contact.email),
        ("Phone", contact.phone),
        ("State", contact.state),
        ("City", contact.city),
        ("Report Date", report_date),
        ("System", "VoltageEnergy Microgrid Advisor v2.0"),
    ]
    for row_index, values in enumerate(rows, start=2):
        _style_table_row(
            ws,
            row_index,
            values,
            fill_color=_alternating_color(row_index),
            height=20,
            font_factory=lambda column: _cfont(11, bold=(column == 1)),
            alignment_factory=lambda _column: Alignment(
                vertical=VERTICAL_CENTER
            ),
        )


def build_excel(req) -> bytes:
    """Build the legacy (non-template) solution Excel report as bytes."""
    ctx = _parse_excel_context(req)
    con = ctx.con
    now = ctx.now
    sim = ctx.sim
    sm = ctx.sm
    ct = ctx.ct

    wb = Workbook()

    # Sheet 1: BOM
    ws1 = wb.active
    ws1.title = "BOM"
    _set_col_w(ws1, [8, 20, 22, 26, 30, 16, 10, 10, 40])

    ws1.row_dimensions[1].height = 58
    ws1.merge_cells("A1:I1")
    _cell_set(
        ws1["A1"],
        "Key Components List / BOM",
        font=_cfont(20, bold=True, color=NAVY),
        align=Alignment(horizontal="center", vertical="center", wrap_text=True),
    )

    ws1.row_dimensions[2].height = 30
    ws1.merge_cells("A2:G2")
    _cell_set(
        ws1["A2"],
        f"Project Name: {con.firstName} {con.lastName} / {con.company}",
        font=_cfont(11, color="2F5496"),
        align=Alignment(horizontal="left", vertical="center"),
        fill_=_sfill(LIGHT_BLUE),
    )
    ws1.merge_cells("H2:I2")
    _cell_set(
        ws1["H2"],
        f"Product Type: {ctx.project_type}",
        font=_cfont(11, color="2F5496"),
        align=Alignment(horizontal="left", vertical="center"),
        fill_=_sfill(LIGHT_BLUE),
    )

    headers = [
        "NO.",
        "Part Number",
        "Part Name",
        "English Name",
        "Specification / Model",
        "Dimensions",
        "Qty",
        "Unit",
        "Notes",
    ]
    ws1.row_dimensions[3].height = 40
    for col_index, header in enumerate(headers, 1):
        cell = ws1.cell(3, col_index, header)
        cell.font = _cfont(11, bold=True, color="FFFFFF")
        cell.fill = _sfill(SUBHDR)
        cell.alignment = Alignment(
            horizontal="center", vertical="center", wrap_text=True
        )
        cell.border = _tborder()

    bom_rows = _build_bom_rows(ctx)
    _write_bom_rows(ws1, bom_rows)

    sig_row = 4 + len(bom_rows) + 1
    ws1.row_dimensions[sig_row].height = 28
    ws1.merge_cells(f"A{sig_row}:I{sig_row}")
    _cell_set(
        ws1[f"A{sig_row}"],
        "Prepared by:              Checked by:              "
        f"Approved by:              Date: {now}",
        font=_cfont(10, color="595959"),
        align=Alignment(horizontal="left", vertical="center"),
        fill_=_sfill(LIGHT_BLUE),
    )

    note_row = sig_row + 1
    ws1.merge_cells(f"A{note_row}:I{note_row}")
    ws1.row_dimensions[note_row].height = 20
    _cell_set(
        ws1[f"A{note_row}"],
        "Notes: All specifications are subject to final purchase-order "
        "confirmation. Quantities should be verified on site.",
        font=_cfont(9, italic=True, color="808080"),
        align=Alignment(horizontal="left", vertical="center"),
    )

    # Sheet 2: Economic analysis / customer summary
    ws2 = wb.create_sheet("Economic Analysis")
    _set_col_w(ws2, [38, 24, 18, 18, 18, 18, 18, 18, 18])

    ws2.merge_cells("A1:B1")
    ws2.row_dimensions[1].height = 26
    _cell_set(
        ws2["A1"],
        "Economic Analysis",
        font=_cfont(13, bold=True, color="FFFFFF"),
        fill_=_sfill(SUBHDR),
        align=Alignment(horizontal="center", vertical="center"),
    )

    current_row = 2

    capex_rows = [
        ("PV Modules", ctx.pv_mod_cost),
        ("Mounting System", ctx.mount_cost),
        ("Battery Storage (BESS)", ctx.bess_cost),
        (DIESEL_GENERATOR_LABEL, ctx.gen_cost),
        ("International Freight", ctx.trans_cost),
        ("Installation", ctx.inst_cost),
        ("Accessories", ctx.acc_cost),
        ("Other Initial Cost", ctx.other_cost),
    ]
    current_row = _write_capex_rows(ws2, current_row, capex_rows)

    pricing_rows = [
        ("Total Cost (excl. margin)", _fmt_money(ctx.subtotal), LIGHT_BLUE),
        (
            f"Gross Profit ({ctx.profit_pct:.1f}%)",
            _fmt_money(ctx.profit_amt),
            WHITE,
        ),
        ("System Quote (incl. margin)", _fmt_money(ctx.selling), ACCENT),
    ]
    current_row = _write_pricing_rows(ws2, current_row, pricing_rows)

    ws2.merge_cells(f"A{current_row}:B{current_row}")
    ws2.row_dimensions[current_row].height = 24
    _cell_set(
        ws2[f"A{current_row}"],
        "Economic Summary",
        font=_cfont(13, bold=True, color="FFFFFF"),
        fill_=_sfill(SUBHDR),
        align=Alignment(horizontal="center", vertical="center"),
    )
    current_row += 1

    econ_rows = [
        ("Annual Load", f"{ctx.ann_load:,.0f} kWh"),
        ("Solar Fraction", f"{float(sim.get('solarFractionPct', 0)):.1f}%"),
        (
            "Annual Fuel Saving",
            f"{float(sim.get('annualFuelSavingLiters', 0)):,.0f} L/yr",
        ),
        (
            "Annual Fuel Cost Saving",
            _fmt_money(float(sim.get("annualFuelSavingUsd", 0))),
        ),
        (
            "Total Cost (excl. margin)",
            _fmt_money(float(sm.get("totalCostUsd", ctx.subtotal))),
        ),
        (
            "System Quote (incl. margin)",
            _fmt_money(float(sm.get("sellingPriceUsd", ctx.selling))),
        ),
        (
            "Gross Profit",
            _fmt_money(float(sm.get("profitAmountUsd", ctx.profit_amt))),
        ),
        ("MG Annual O&M", _fmt_money(float(sm.get("mgAnnualOmUsd", 0)))),
        (
            "MG Annual Fuel Cost",
            _fmt_money(float(sm.get("mgAnnualFuelUsd", 0))),
        ),
        (
            "Diesel-only Annual Fuel Cost",
            _fmt_money(float(sm.get("dieselAnnualFuelUsd", 0))),
        ),
        ("Microgrid LCOE", f"$ {float(sm.get('finalMgLcoe', 0)):.4f}/kWh"),
        (
            "Diesel-only LCOE",
            f"$ {float(sm.get('finalDieselLcoe', 0)):.4f}/kWh",
        ),
        ("Payback Period", f"Year {sm.get('breakevenYear', '-')}"),
        (
            "20-year Net Revenue",
            _fmt_money(float(sm.get("finalCumulativeRevenue", 0))),
        ),
    ]
    current_row = _write_economic_rows(ws2, current_row, econ_rows)

    current_row = _write_site_layout_note(
        ws2,
        current_row,
        ctx.layout_note,
        ctx.measured_site_area,
        ctx.usable_site_area,
        ctx.layout_max_sets,
    )

    current_row = _write_comparison_table(
        ws2, current_row, ct, sm.get("breakevenYear")
    )

    _add_contact_sheet(wb, con, now)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()

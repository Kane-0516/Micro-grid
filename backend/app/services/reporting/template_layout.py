"""Excel 报告模板的单元格布局与默认样式常量."""

from __future__ import annotations

import os
from dataclasses import dataclass
from functools import lru_cache
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

BACKEND_ROOT = Path(__file__).resolve().parents[3]
PRIVATE_TEMPLATE_DIR = BACKEND_ROOT / ".private" / "report_templates"


def _resolve_template_path(env_var: str, file_name: str) -> Path:
    raw_path = os.getenv(env_var, "").strip()
    if raw_path:
        return Path(raw_path).expanduser()
    return PRIVATE_TEMPLATE_DIR / file_name


INTERNAL_TEMPLATE_PATH = _resolve_template_path(
    "MICROGRID_INTERNAL_TEMPLATE_PATH",
    "\u516c\u53f8\u5185\u90e8\u6210\u672c\u7edf\u8ba1.xlsx",
)
CUSTOMER_TEMPLATE_PATH = _resolve_template_path(
    "MICROGRID_CUSTOMER_TEMPLATE_PATH",
    "\u987e\u5ba2\u6210\u672c\u7edf\u8ba1.xlsx",
)

YES = "\u662f / Yes"
NO = "\u5426 / No"
USD = "\u7f8e\u5143 / USD"
PV_SYSTEM_LABEL = "\u5149\u4f0f\u7cfb\u7edf / PV System"
BESS_SYSTEM_LABEL = "\u50a8\u80fd\u7cfb\u7edf / BESS System"
DIESEL_SYSTEM_LABEL = "\u67f4\u6cb9\u673a\u7cfb\u7edf / Diesel System"
ELECTRICAL_SYSTEM_LABEL = "\u7535\u6c14\u914d\u5957 / Electrical Balance"
LOGISTICS_SYSTEM_LABEL = (
    "\u8fd0\u8f93\u5b89\u88c5\u5176\u4ed6 / Logistics & Installation"
)
CONNECTORS_LABEL = "\u8fde\u63a5\u5668/\u8f85\u6750 / Connectors-Auxiliary"
TANK_ATS_LABEL = "\u6cb9\u7bb1/ATS / Tank-ATS"
OTHER_LABEL = "\u5176\u4ed6 / Other"
ANNUAL_FIXED_OM_LABEL = (
    "\u5e74\u5ea6\u56fa\u5b9a\u8fd0\u7ef4 / Annual Fixed O&M"
)

INTERNAL_SHEET_INDEX = {
    "instructions": 0,
    "basic": 1,
    "capex": 2,
    "pv": 3,
    "bess": 4,
    "diesel": 5,
    "electrical": 6,
    "logistics": 7,
    "opex": 8,
    "fuel": 9,
    "diesel_maintenance": 10,
    "dashboard": 11,
}

CUSTOMER_SHEET_INDEX = {
    "basic": 0,
    "capex": 1,
    "pv": 2,
    "bess": 3,
    "diesel": 4,
    "electrical": 5,
}
SHARED_TEMPLATE_SECTIONS = ("pv", "bess", "diesel", "electrical")
DETAIL_LABEL_COLUMNS = ("B", "C", "D")


@dataclass(frozen=True)
class TemplateBasicLayout:
    """Cell layout for the template's basic-info input rows."""

    sheet_key: str
    input_column: str
    rows: dict[str, int]


@dataclass(frozen=True)
class TemplateCapexLayout:
    """Cell layout for the template's CAPEX summary rows."""

    sheet_key: str
    rows: dict[str, int]
    amount_column: str | None = None


@dataclass(frozen=True)
class TemplateSectionLayout:
    """Cell layout for one detail section (e.g. PV, BESS, diesel)."""

    sheet_key: str
    header_row: int
    rows: dict[str, int]
    total_row: int | None = None


@dataclass(frozen=True)
class WorkbookLayout:
    """Full cell layout for a report workbook template."""

    basic: TemplateBasicLayout
    capex: TemplateCapexLayout
    sections: dict[str, TemplateSectionLayout]
    fuel_input_cells: dict[str, str] | None = None
    diesel_maintenance_adjustment_cells: dict[str, str] | None = None


@dataclass(frozen=True)
class TemplateDetailRowDefault:
    """Default values for one CAPEX detail line row."""

    key: str
    row: int
    system: str
    subcategory: str
    item: str
    brand: str
    model: str
    spec: str
    qty_unit: str
    qty: Any
    unit_price: Any
    currency: str
    fx_to_usd: Any
    tax_rate: Any
    freight_usd: Any
    other_adjust_usd: Any
    source: str
    remarks: str
    include: bool


@dataclass(frozen=True)
class TemplateOpexRowDefault:
    """Default values for one OPEX line row."""

    key: str
    row: int
    category: str
    subcategory: str
    item: str
    basis: str
    qty_unit: str
    qty: Any
    unit_price: Any
    currency: str
    fx_to_usd: Any
    source: str
    remarks: str
    include: bool


@dataclass(frozen=True)
class InternalTemplateDefaults:
    """Default detail/opex rows for every section of the template."""

    pv: dict[str, TemplateDetailRowDefault]
    bess: dict[str, TemplateDetailRowDefault]
    diesel: dict[str, TemplateDetailRowDefault]
    electrical: dict[str, TemplateDetailRowDefault]
    logistics: dict[str, TemplateDetailRowDefault]
    opex: dict[str, TemplateOpexRowDefault]


BASIC_ROWS = {
    "project_name": 5,
    "client_owner": 6,
    "country_region": 7,
    "site": 8,
    "quotation_date": 9,
    "project_stage": 10,
    "pv_size_kwp": 14,
    "bess_size_kwh": 15,
    "diesel_size_kw": 16,
    "daily_energy_demand_kwh": 17,
    "design_life_years": 18,
    "base_currency": 19,
    "tax_basis_notes": 20,
    "project_notes": 21,
}

BASIC_FIELD_LABELS = {
    "project_name": "\u9879\u76ee\u540d\u79f0 / Project Name",
    "client_owner": "\u5ba2\u6237/\u4e1a\u4e3b / Client-Owner",
    "country_region": "\u56fd\u5bb6/\u5730\u533a / Country-Region",
    "site": "\u5177\u4f53\u7ad9\u70b9 / Site",
    "quotation_date": "\u62a5\u4ef7\u65e5\u671f / Quotation Date",
    "project_stage": "\u9879\u76ee\u9636\u6bb5 / Project Stage",
    "pv_size_kwp": "\u5149\u4f0f\u89c4\u6a21 / PV Size",
    "bess_size_kwh": "\u50a8\u80fd\u89c4\u6a21 / BESS Size",
    "diesel_size_kw": "\u67f4\u6cb9\u673a\u5bb9\u91cf / Diesel Size",
    "daily_energy_demand_kwh": "\u8d1f\u8377\u65e5\u5747\u7528\u7535\u91cf / "
    "Daily Energy Demand",
    "design_life_years": "\u8bbe\u8ba1\u5e74\u9650 / Design Life",
    "base_currency": "\u57fa\u51c6\u5e01\u79cd / Base Currency",
    "tax_basis_notes": "\u7a0e\u8d39\u53e3\u5f84\u8bf4\u660e / Tax Basis Notes",
    "project_notes": "\u9879\u76ee\u5907\u6ce8 / Project Notes",
}

CAPEX_LABELS = {
    "pv": PV_SYSTEM_LABEL,
    "bess": BESS_SYSTEM_LABEL,
    "diesel": DIESEL_SYSTEM_LABEL,
    "electrical": ELECTRICAL_SYSTEM_LABEL,
    "logistics": LOGISTICS_SYSTEM_LABEL,
    "total": "CAPEX \u603b\u8ba1 / Total CAPEX",
}

SECTION_LABELS: dict[str, dict[str, tuple[str, str, str]]] = {
    "pv": {
        "module_primary": (
            PV_SYSTEM_LABEL,
            "\u7ec4\u4ef6 / Module",
            "\u5149\u4f0f\u7ec4\u4ef6 / PV Module",
        ),
        "module_secondary": (
            PV_SYSTEM_LABEL,
            "\u7ec4\u4ef6 / Module",
            "\u5149\u4f0f\u7ec4\u4ef6 / PV Module",
        ),
        "inverter": (
            PV_SYSTEM_LABEL,
            "\u9006\u53d8\u5668 / Inverter",
            "\u7ec4\u4e32\u9006\u53d8\u5668 / String Inverter",
        ),
        "combiner": (
            PV_SYSTEM_LABEL,
            "\u9006\u53d8\u5668 / Inverter",
            "\u6c47\u6d41\u7bb1 / Combine",
        ),
        "mounting": (
            PV_SYSTEM_LABEL,
            "\u652f\u67b6 / Mounting",
            "\u5149\u4f0f\u652f\u67b6 / PV Mounting",
        ),
        "mc4": (
            PV_SYSTEM_LABEL,
            CONNECTORS_LABEL,
            "\u8fde\u63a5\u5668/MC4 / Connectors-MC4",
        ),
        "dc_combiner_box": (
            PV_SYSTEM_LABEL,
            CONNECTORS_LABEL,
            "\u76f4\u6d41\u6c47\u6d41/\u7ebf\u76d2 / DC Combiner-Junction Box",
        ),
        "other_aux": (
            PV_SYSTEM_LABEL,
            CONNECTORS_LABEL,
            "\u5176\u4ed6\u5149\u4f0f\u8f85\u6750 / Other PV Auxiliary",
        ),
    },
    "bess": {
        "pcs": (
            BESS_SYSTEM_LABEL,
            "PCS / PCS",
            "\u50a8\u80fd\u53d8\u6d41\u5668 PCS / PCS",
        ),
        "battery_pack": (
            BESS_SYSTEM_LABEL,
            "\u7535\u6c60\u5305 / Battery Pack",
            "\u7535\u6c60\u5305/\u6a21\u7ec4 / Battery Pack-Module",
        ),
        "integrated_pallet": (
            BESS_SYSTEM_LABEL,
            "\u5149\u50a8\u4e00\u4f53\u5316\u6258\u76d8/PV-ESS Integrated "
            "Pallet",
            "\u4e00\u4f53\u5316\u6258\u76d8/Integrated pallet",
        ),
        "battery_cabinet": (
            BESS_SYSTEM_LABEL,
            "\u7535\u6c60\u67dc/\u96c6\u88c5\u7bb1 / Battery Cabinet-Container",
            "\u7535\u6c60\u67dc/\u7535\u6c60\u7c07 / Battery Cabinet-String",
        ),
        "bms": (BESS_SYSTEM_LABEL, "BMS/EMS / BMS-EMS", "BMS / BMS"),
        "ems": (
            BESS_SYSTEM_LABEL,
            "BMS/EMS / BMS-EMS",
            "EMS/\u80fd\u91cf\u7ba1\u7406 / EMS-Energy Management",
        ),
        "hvac": (
            BESS_SYSTEM_LABEL,
            "HVAC/\u6d88\u9632 / HVAC-Fire",
            "\u7a7a\u8c03/HVAC / HVAC",
        ),
        "fire_suppression": (
            BESS_SYSTEM_LABEL,
            "HVAC/\u6d88\u9632 / HVAC-Fire",
            "\u6d88\u9632\u7cfb\u7edf / Fire Suppression",
        ),
        "bess_aux": (
            BESS_SYSTEM_LABEL,
            CONNECTORS_LABEL,
            "\u50a8\u80fd\u8f85\u6750/\u94dc\u6392 / BESS Auxiliary-Busbar",
        ),
    },
    "diesel": {
        "generator": (
            DIESEL_SYSTEM_LABEL,
            "\u67f4\u6cb9\u53d1\u7535\u673a / Diesel Generator",
            "\u67f4\u6cb9\u53d1\u7535\u673a\u7ec4 / Diesel Generator Set",
        ),
        "ats_panel": (
            DIESEL_SYSTEM_LABEL,
            TANK_ATS_LABEL,
            "ATS/\u5e76\u673a\u67dc / ATS-Paralleling Panel",
        ),
        "day_tank": (
            DIESEL_SYSTEM_LABEL,
            TANK_ATS_LABEL,
            "\u65e5\u7528\u6cb9\u7bb1/\u5e95\u5ea7\u6cb9\u7bb1 / Day Tank-Base "
            "Tank",
        ),
        "exhaust": (
            DIESEL_SYSTEM_LABEL,
            TANK_ATS_LABEL,
            "\u6392\u70df/\u6d88\u97f3\u7cfb\u7edf / Exhaust-Silencer",
        ),
        "starter_battery": (
            DIESEL_SYSTEM_LABEL,
            CONNECTORS_LABEL,
            "\u542f\u52a8\u7535\u6c60/\u5145\u7535\u5668 / Starter "
            "Battery-Charger",
        ),
        "diesel_aux": (
            DIESEL_SYSTEM_LABEL,
            CONNECTORS_LABEL,
            "\u67f4\u6cb9\u673a\u8f85\u6750 / Diesel Auxiliary",
        ),
    },
    "electrical": {
        "switchgear": (
            ELECTRICAL_SYSTEM_LABEL,
            "\u5f00\u5173\u67dc / Switchgear",
            "\u4f4e\u538b\u5f00\u5173\u67dc/\u914d\u7535\u67dc / LV "
            "Switchboard-Distribution Panel",
        ),
        "transformer": (
            ELECTRICAL_SYSTEM_LABEL,
            "\u53d8\u538b\u5668 / Transformer",
            "\u53d8\u538b\u5668 / Transformer",
        ),
        "ac_cable": (
            ELECTRICAL_SYSTEM_LABEL,
            "AC\u7535\u7f06 / AC Cable",
            "\u4ea4\u6d41\u7535\u7f06 / AC Cable",
        ),
        "dc_cable": (
            ELECTRICAL_SYSTEM_LABEL,
            "DC\u7535\u7f06 / DC Cable",
            "\u76f4\u6d41\u7535\u7f06 / DC Cable",
        ),
        "grounding": (
            ELECTRICAL_SYSTEM_LABEL,
            "\u63a5\u5730\u9632\u96f7 / Grounding-Lightning",
            "\u63a5\u5730\u4e0e\u9632\u96f7 / Grounding & Lightning",
        ),
        "scada": (
            ELECTRICAL_SYSTEM_LABEL,
            "\u901a\u4fe1\u76d1\u63a7 / Communication-SCADA",
            "\u901a\u4fe1\u76d1\u63a7/SCADA / Communication-SCADA",
        ),
        "cable_tray": (
            ELECTRICAL_SYSTEM_LABEL,
            CONNECTORS_LABEL,
            "\u6865\u67b6/\u7ebf\u69fd/\u7aef\u5b50 / Cable "
            "Tray-Trunking-Terminals",
        ),
        "electrical_aux": (
            ELECTRICAL_SYSTEM_LABEL,
            CONNECTORS_LABEL,
            "\u5176\u4ed6\u7535\u6c14\u914d\u5957 / Other Electrical Balance",
        ),
    },
    "logistics": {
        "international_transport": (
            LOGISTICS_SYSTEM_LABEL,
            "\u8fd0\u8f93 / Transportation",
            "\u56fd\u9645\u8fd0\u8f93 / International Transportation",
        ),
        "domestic_transport": (
            LOGISTICS_SYSTEM_LABEL,
            "\u8fd0\u8f93 / Transportation",
            "\u56fd\u5185/\u5185\u9646\u8fd0\u8f93 / Domestic-Inland "
            "Transportation",
        ),
        "installation_labor": (
            LOGISTICS_SYSTEM_LABEL,
            "\u5b89\u88c5 / Installation",
            "\u5b89\u88c5\u4eba\u5de5 / Installation Labor",
        ),
        "commissioning": (
            LOGISTICS_SYSTEM_LABEL,
            "\u8c03\u8bd5 / Commissioning",
            "\u8c03\u8bd5/\u8bd5\u8fd0\u884c / Commissioning-Test Run",
        ),
        "crane_tools": (
            LOGISTICS_SYSTEM_LABEL,
            "\u5b89\u88c5 / Installation",
            "\u540a\u88c5/\u53c9\u8f66/\u673a\u5177 / Crane-Forklift-Tools",
        ),
        "customs_clearance": (
            LOGISTICS_SYSTEM_LABEL,
            "\u5173\u7a0e/\u6e05\u5173 /Customs Clearance",
            "\u6e05\u5173/\u62a5\u5173 / Customs Clearance",
        ),
        "travel_accommodation": (
            LOGISTICS_SYSTEM_LABEL,
            OTHER_LABEL,
            "\u5dee\u65c5/\u4f4f\u5bbf / Travel-Accommodation",
        ),
        "contingency": (
            LOGISTICS_SYSTEM_LABEL,
            OTHER_LABEL,
            "\u9884\u5907\u8d39/\u5176\u4ed6\u8c03\u6574 / Contingency-Other "
            "Adjustment",
        ),
    },
    "opex": {
        "microgrid_om": (
            ANNUAL_FIXED_OM_LABEL,
            "\u8fd0\u7ef4 / O&M",
            "\u5fae\u7535\u7f51\u8fd0\u7ef4\u8d39\u7528 / Microgrid O&M",
        ),
        "equipment_repair": (
            ANNUAL_FIXED_OM_LABEL,
            "\u7ef4\u4fee / Repair",
            "\u8bbe\u5907\u7ef4\u4fee\u8d39\u7528 / Equipment Repair",
        ),
        "insurance": (
            ANNUAL_FIXED_OM_LABEL,
            "\u4fdd\u9669 / Insurance",
            "\u4fdd\u9669\u8d39\u7528 / Insurance",
        ),
        "diesel_maintenance": (
            ANNUAL_FIXED_OM_LABEL,
            "\u7ef4\u4fdd / Maintenance",
            "\u67f4\u6cb9\u673a\u5e74\u5ea6\u4fdd\u517b / Annual Diesel "
            "Maintenance",
        ),
        "remote_monitoring": (
            ANNUAL_FIXED_OM_LABEL,
            "\u901a\u4fe1 / Communication",
            "\u8fdc\u7a0b\u76d1\u63a7/\u901a\u4fe1 / Remote "
            "Monitoring-Communication",
        ),
        "site_lease": (
            ANNUAL_FIXED_OM_LABEL,
            "\u79df\u8d41/\u571f\u5730 / Lease-Land",
            "\u573a\u5730\u79df\u8d41/\u5730\u79df / Site Lease-Land Rent",
        ),
        "other_annual": (
            ANNUAL_FIXED_OM_LABEL,
            OTHER_LABEL,
            "\u5176\u4ed6\u5e74\u5ea6\u8d39\u7528 / Other Annual Cost",
        ),
    },
}


def _section(
    sheet_key: str, header_row: int, total_row: int | None, rows: dict[str, int]
) -> TemplateSectionLayout:
    return TemplateSectionLayout(
        sheet_key=sheet_key,
        header_row=header_row,
        total_row=total_row,
        rows=rows,
    )


@lru_cache(maxsize=1)
def get_internal_template_layout() -> WorkbookLayout:
    """Return the internal (full-detail) report template's cell layout."""
    return WorkbookLayout(
        basic=TemplateBasicLayout("basic", "B", BASIC_ROWS),
        capex=TemplateCapexLayout(
            "capex",
            {
                "pv": 5,
                "bess": 6,
                "diesel": 7,
                "electrical": 8,
                "logistics": 9,
                "total": 11,
            },
            "D",
        ),
        sections={
            "pv": _section(
                "pv",
                4,
                13,
                {
                    "module_primary": 5,
                    "module_secondary": 6,
                    "inverter": 7,
                    "combiner": 8,
                    "mounting": 9,
                    "mc4": 10,
                    "dc_combiner_box": 11,
                    "other_aux": 12,
                },
            ),
            "bess": _section(
                "bess",
                4,
                14,
                {
                    "pcs": 5,
                    "battery_pack": 6,
                    "integrated_pallet": 7,
                    "battery_cabinet": 8,
                    "bms": 9,
                    "ems": 10,
                    "hvac": 11,
                    "fire_suppression": 12,
                    "bess_aux": 13,
                },
            ),
            "diesel": _section(
                "diesel",
                4,
                11,
                {
                    "generator": 5,
                    "ats_panel": 6,
                    "day_tank": 7,
                    "exhaust": 8,
                    "starter_battery": 9,
                    "diesel_aux": 10,
                },
            ),
            "electrical": _section(
                "electrical",
                4,
                13,
                {
                    "switchgear": 5,
                    "transformer": 6,
                    "ac_cable": 7,
                    "dc_cable": 8,
                    "grounding": 9,
                    "scada": 10,
                    "cable_tray": 11,
                    "electrical_aux": 12,
                },
            ),
            "logistics": _section(
                "logistics",
                4,
                13,
                {
                    "international_transport": 5,
                    "domestic_transport": 6,
                    "installation_labor": 7,
                    "commissioning": 8,
                    "crane_tools": 9,
                    "customs_clearance": 10,
                    "travel_accommodation": 11,
                    "contingency": 12,
                },
            ),
            "opex": _section(
                "opex",
                4,
                12,
                {
                    "microgrid_om": 5,
                    "equipment_repair": 6,
                    "insurance": 7,
                    "diesel_maintenance": 8,
                    "remote_monitoring": 9,
                    "site_lease": 10,
                    "other_annual": 11,
                },
            ),
        },
        fuel_input_cells={
            "diesel_price_usd_per_liter": "E5",
            "delivered_fuel_adder_usd_per_liter": "E6",
            "estimated_annual_runtime_hours": "E7",
            "average_diesel_load_factor_pct": "E8",
            "estimated_annual_fuel_use_liters": "E9",
            "fuel_price_upside_pct": "E10",
        },
        diesel_maintenance_adjustment_cells={
            "S_oil": "H5",
            "t_pm": "H6",
            "V_oil": "H7",
            "P_oil": "H8",
            "C_of": "H9",
            "C_ff": "H10",
            "C_af": "H11",
            "labor_rate": "H12",
            "S_air": "H13",
            "coolant_interval_years": "H14",
            "coolant_gal_per_service": "H15",
            "P_coolant": "H16",
            "coolant_labor_hours": "H17",
            "battery_interval_years": "H18",
            "C_battery": "H19",
            "battery_labor_hours": "H20",
        },
    )


@lru_cache(maxsize=1)
def get_customer_template_layout() -> WorkbookLayout:
    """Return the customer-facing report template's cell layout."""
    return WorkbookLayout(
        basic=TemplateBasicLayout("basic", "B", BASIC_ROWS),
        capex=TemplateCapexLayout(
            "capex",
            {
                "pv": 4,
                "bess": 5,
                "diesel": 6,
                "electrical": 7,
                "logistics": 8,
                "total": 10,
            },
        ),
        sections={
            "pv": _section(
                "pv",
                3,
                None,
                {
                    "module_primary": 4,
                    "module_secondary": 5,
                    "inverter": 6,
                    "combiner": 7,
                    "mounting": 8,
                    "mc4": 9,
                    "dc_combiner_box": 10,
                    "other_aux": 11,
                },
            ),
            "bess": _section(
                "bess",
                3,
                None,
                {
                    "pcs": 4,
                    "battery_pack": 5,
                    "integrated_pallet": 6,
                    "battery_cabinet": 7,
                    "bms": 8,
                    "ems": 9,
                    "hvac": 10,
                    "fire_suppression": 11,
                    "bess_aux": 12,
                },
            ),
            "diesel": _section(
                "diesel",
                3,
                None,
                {
                    "generator": 4,
                    "ats_panel": 5,
                    "day_tank": 6,
                    "exhaust": 7,
                    "starter_battery": 8,
                    "diesel_aux": 9,
                },
            ),
            "electrical": _section(
                "electrical",
                3,
                None,
                {
                    "switchgear": 4,
                    "transformer": 5,
                    "ac_cable": 6,
                    "dc_cable": 7,
                    "grounding": 8,
                    "scada": 9,
                    "cable_tray": 10,
                    "electrical_aux": 11,
                },
            ),
        },
    )


def internal_sheet(workbook, key: str):
    """Return the internal template's worksheet for the given section key."""
    return workbook[workbook.sheetnames[INTERNAL_SHEET_INDEX[key]]]


def customer_sheet(workbook, key: str):
    """Return the customer template's worksheet for the given section key."""
    return workbook[workbook.sheetnames[CUSTOMER_SHEET_INDEX[key]]]


def _as_text(value: Any) -> str:
    return "" if value in (None, "") else str(value)


def _label_triplet(ws, row: int) -> tuple[str, str, str]:
    return tuple(
        _as_text(ws[f"{column}{row}"].value) for column in DETAIL_LABEL_COLUMNS
    )


def _read_detail_defaults(
    ws, rows_by_key: dict[str, int]
) -> dict[str, TemplateDetailRowDefault]:
    return {
        key: TemplateDetailRowDefault(
            key=key,
            row=row,
            system=_as_text(ws[f"B{row}"].value),
            subcategory=_as_text(ws[f"C{row}"].value),
            item=_as_text(ws[f"D{row}"].value),
            brand=_as_text(ws[f"E{row}"].value),
            model=_as_text(ws[f"F{row}"].value),
            spec=_as_text(ws[f"G{row}"].value),
            qty_unit=_as_text(ws[f"H{row}"].value),
            qty=ws[f"I{row}"].value,
            unit_price=ws[f"J{row}"].value,
            currency=_as_text(ws[f"K{row}"].value),
            fx_to_usd=ws[f"L{row}"].value,
            tax_rate=ws[f"M{row}"].value,
            freight_usd=ws[f"N{row}"].value,
            other_adjust_usd=ws[f"O{row}"].value,
            source=_as_text(ws[f"Q{row}"].value),
            remarks=_as_text(ws[f"R{row}"].value),
            include=_as_text(ws[f"S{row}"].value) == YES,
        )
        for key, row in rows_by_key.items()
    }


def _read_opex_defaults(
    ws, rows_by_key: dict[str, int]
) -> dict[str, TemplateOpexRowDefault]:
    return {
        key: TemplateOpexRowDefault(
            key=key,
            row=row,
            category=_as_text(ws[f"B{row}"].value),
            subcategory=_as_text(ws[f"C{row}"].value),
            item=_as_text(ws[f"D{row}"].value),
            basis=_as_text(ws[f"E{row}"].value),
            qty_unit=_as_text(ws[f"F{row}"].value),
            qty=ws[f"G{row}"].value,
            unit_price=ws[f"H{row}"].value,
            currency=_as_text(ws[f"I{row}"].value),
            fx_to_usd=ws[f"J{row}"].value,
            source=_as_text(ws[f"L{row}"].value),
            remarks=_as_text(ws[f"M{row}"].value),
            include=_as_text(ws[f"N{row}"].value) == YES,
        )
        for key, row in rows_by_key.items()
    }


@lru_cache(maxsize=1)
def get_internal_template_defaults() -> InternalTemplateDefaults:
    """Read the internal template's default detail/opex row values."""
    layout = get_internal_template_layout()
    wb = load_workbook(INTERNAL_TEMPLATE_PATH, data_only=False)
    return InternalTemplateDefaults(
        pv=_read_detail_defaults(
            internal_sheet(wb, "pv"), layout.sections["pv"].rows
        ),
        bess=_read_detail_defaults(
            internal_sheet(wb, "bess"), layout.sections["bess"].rows
        ),
        diesel=_read_detail_defaults(
            internal_sheet(wb, "diesel"), layout.sections["diesel"].rows
        ),
        electrical=_read_detail_defaults(
            internal_sheet(wb, "electrical"), layout.sections["electrical"].rows
        ),
        logistics=_read_detail_defaults(
            internal_sheet(wb, "logistics"), layout.sections["logistics"].rows
        ),
        opex=_read_opex_defaults(
            internal_sheet(wb, "opex"), layout.sections["opex"].rows
        ),
    )


def _validate_basic(
    ws, layout: TemplateBasicLayout, mismatches: list[str], prefix: str
) -> None:
    for key, row in layout.rows.items():
        if _as_text(ws[f"A{row}"].value) != BASIC_FIELD_LABELS[key]:
            mismatches.append(f"{prefix}:{key}@A{row}")


def _validate_capex(
    ws, layout: TemplateCapexLayout, mismatches: list[str], prefix: str
) -> None:
    for key, row in layout.rows.items():
        col = "A" if key == "total" else "B"
        if _as_text(ws[f"{col}{row}"].value) != CAPEX_LABELS[key]:
            mismatches.append(f"{prefix}:{key}@{col}{row}")


def _validate_section(
    ws,
    layout: TemplateSectionLayout,
    section_key: str,
    mismatches: list[str],
    prefix: str,
) -> None:
    labels = SECTION_LABELS[section_key]
    for key, row in layout.rows.items():
        if _label_triplet(ws, row) != labels[key]:
            mismatches.append(f"{prefix}:{key}@{row}")


def _validate_workbook_layout(
    workbook,
    layout: WorkbookLayout,
    sheet_getter,
    prefix: str,
    mismatches: list[str],
    section_keys: tuple[str, ...],
) -> None:
    _validate_basic(
        sheet_getter(workbook, "basic"),
        layout.basic,
        mismatches,
        f"{prefix}.basic",
    )
    _validate_capex(
        sheet_getter(workbook, "capex"),
        layout.capex,
        mismatches,
        f"{prefix}.capex",
    )
    for section_key in section_keys:
        _validate_section(
            sheet_getter(workbook, section_key),
            layout.sections[section_key],
            section_key,
            mismatches,
            f"{prefix}.{section_key}",
        )


@lru_cache(maxsize=1)
def validate_template_layouts() -> None:
    """Validate that both template workbooks match their declared layouts."""
    internal_layout = get_internal_template_layout()
    customer_layout = get_customer_template_layout()
    internal_wb = load_workbook(INTERNAL_TEMPLATE_PATH, data_only=False)
    customer_wb = load_workbook(CUSTOMER_TEMPLATE_PATH, data_only=False)
    mismatches: list[str] = []

    _validate_workbook_layout(
        internal_wb,
        internal_layout,
        internal_sheet,
        "internal",
        mismatches,
        SHARED_TEMPLATE_SECTIONS + ("logistics", "opex"),
    )
    _validate_workbook_layout(
        customer_wb,
        customer_layout,
        customer_sheet,
        "customer",
        mismatches,
        SHARED_TEMPLATE_SECTIONS,
    )

    if mismatches:
        raise ValueError(
            "Template row layout changed and no longer matches backend "
            "mappings: " + ", ".join(mismatches)
        )


DETAIL_SUBTOTAL_FORMULA = '=IF(OR(I5="",J5=""),"",I5*J5*IF(L5="",1,L5)*(1+IF(M5="",0,M5))+IF(N5="",0,N5)+IF(O5="",0,O5))'  # noqa: E501
EIGHT_ROW_INCLUDED_SUM_FORMULA = '=SUMIFS($P$5:$P$12,$S$5:$S$12,"\u662f / Yes")'

EXPECTED_INTERNAL_FORMULAS = {
    (
        "capex",
        "D5",
    ): "=IF('03_\u5149\u4f0f\u7cfb\u7edf\u660e\u7ec6'!P13=\"\",'','03_\u5149\u4f0f\u7cfb\u7edf\u660e\u7ec6'!P13)",  # noqa: E501
    (
        "capex",
        "D6",
    ): "=IF('04_\u50a8\u80fd\u7cfb\u7edf\u660e\u7ec6'!P14=\"\",'','04_\u50a8\u80fd\u7cfb\u7edf\u660e\u7ec6'!P14)",  # noqa: E501
    (
        "capex",
        "D7",
    ): "=IF('05_\u67f4\u6cb9\u673a\u7cfb\u7edf\u660e\u7ec6'!P11=\"\",'','05_\u67f4\u6cb9\u673a\u7cfb\u7edf\u660e\u7ec6'!P11)",  # noqa: E501
    (
        "capex",
        "D8",
    ): "=IF('06_\u7535\u6c14\u914d\u5957\u660e\u7ec6'!P13=\"\",'','06_\u7535\u6c14\u914d\u5957\u660e\u7ec6'!P13)",  # noqa: E501
    (
        "capex",
        "D9",
    ): "=IF('07_\u8fd0\u8f93\u5b89\u88c5\u5176\u4ed6'!P13=\"\",'','07_\u8fd0\u8f93\u5b89\u88c5\u5176\u4ed6'!P13)",  # noqa: E501
    ("capex", "D11"): "=SUM(D5:D9)",
    ("pv", "P5"): DETAIL_SUBTOTAL_FORMULA,
    ("pv", "P13"): EIGHT_ROW_INCLUDED_SUM_FORMULA,
    ("bess", "P5"): DETAIL_SUBTOTAL_FORMULA,
    ("bess", "P14"): '=SUMIFS($P$5:$P$13,$S$5:$S$13,"\u662f / Yes")',
    ("diesel", "P5"): DETAIL_SUBTOTAL_FORMULA,
    ("diesel", "P11"): '=SUMIFS($P$5:$P$10,$S$5:$S$10,"\u662f / Yes")',
    ("electrical", "P5"): DETAIL_SUBTOTAL_FORMULA,
    ("electrical", "P13"): EIGHT_ROW_INCLUDED_SUM_FORMULA,
    ("logistics", "P5"): DETAIL_SUBTOTAL_FORMULA,
    ("logistics", "P13"): EIGHT_ROW_INCLUDED_SUM_FORMULA,
    (
        "opex",
        "H8",
    ): "=IF('10_\u67f4\u6cb9\u673a\u7ef4\u4fdd\u53c2\u6570'!E31=\"\",\"\",'10_\u67f4\u6cb9\u673a\u7ef4\u4fdd\u53c2\u6570'!E31)",  # noqa: E501
    ("opex", "K5"): '=IF(OR(G5="",H5=""),"",G5*H5*IF(J5="",1,J5))',
    ("opex", "K12"): '=SUMIFS(K5:K11,N5:N11,"\u662f / Yes")',
    (
        "diesel_maintenance",
        "E25",
    ): "=IF('09_\u67f4\u6cb9\u4e0e\u80fd\u6e90\u53c2\u6570'!E7=\"\",\"\",'09_\u67f4\u6cb9\u4e0e\u80fd\u6e90\u53c2\u6570'!E7)",  # noqa: E501
    (
        "diesel_maintenance",
        "E26",
    ): '=IF(E25="","",ROUNDUP(E25/IF($H$5<>"",$H$5,$E$5),0))',
    (
        "diesel_maintenance",
        "E27",
    ): '=IF(E26="","",E26*((IF($H$6<>"",$H$6,$E$6)*IF($H$12<>"",$H$12,$E$12))+(IF($H$7<>"",$H$7,$E$7)*IF($H$8<>"",$H$8,$E$8))+IF($H$9<>"",$H$9,$E$9)+IF($H$10<>"",$H$10,$E$10)))',  # noqa: E501
    (
        "diesel_maintenance",
        "E28",
    ): '=IF(E25="","",IF(IF($H$13<>"",$H$13,$E$13)="",ROUNDUP(E26/2,0)*IF($H$11<>"",$H$11,$E$11),ROUNDUP(E25/IF($H$13<>"",$H$13,$E$13),0)*IF($H$11<>"",$H$11,$E$11)))',  # noqa: E501
    (
        "diesel_maintenance",
        "E29",
    ): '=IF(OR(IF($H$14<>"",$H$14,$E$14)="",IF($H$14<>"",$H$14,$E$14)<=0),"",((IF($H$15<>"",$H$15,$E$15)*IF($H$16<>"",$H$16,$E$16))+(IF($H$17<>"",$H$17,$E$17)*IF($H$12<>"",$H$12,$E$12)))/IF($H$14<>"",$H$14,$E$14))',  # noqa: E501
    (
        "diesel_maintenance",
        "E30",
    ): '=IF(OR(IF($H$18<>"",$H$18,$E$18)="",IF($H$18<>"",$H$18,$E$18)<=0),"",(IF($H$19<>"",$H$19,$E$19)+(IF($H$20<>"",$H$20,$E$20)*IF($H$12<>"",$H$12,$E$12)))/IF($H$18<>"",$H$18,$E$18))',  # noqa: E501
    ("diesel_maintenance", "E31"): "=SUM(E27:E30)",
}


def _normalize_formula(value: Any) -> str:
    if value is None:
        return ""
    return str(value).replace(" ", "").replace('""', "''")


@lru_cache(maxsize=1)
def validate_internal_template_formulas() -> None:
    """Validate that key internal-template formula cells are unchanged."""
    wb = load_workbook(INTERNAL_TEMPLATE_PATH, data_only=False)
    mismatches: list[str] = []
    for (sheet_key, cell), expected in EXPECTED_INTERNAL_FORMULAS.items():
        actual = internal_sheet(wb, sheet_key)[cell].value
        if _normalize_formula(actual) != _normalize_formula(expected):
            mismatches.append(f"{sheet_key}!{cell}")
    if mismatches:
        raise ValueError(
            "Internal template formulas changed and no longer match backend "
            "mappings: " + ", ".join(mismatches)
        )
